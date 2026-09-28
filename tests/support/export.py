"""Call /v1/export and parse the three formats back into record-id lists."""

import io
import json
import re

import pdfplumber
from openpyxl import load_workbook

UUID = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b")


def export(api, auth, persona: str, fmt: str = "json", **body):
    return api.post("/v1/export", json={"format": fmt, **body}, headers=auth(persona))


def json_rows(content: bytes) -> tuple[dict, list[dict]]:
    doc = json.loads(content)
    return doc["metadata"], doc["records"]


def xlsx_rows(content: bytes) -> tuple[dict, list[dict]]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(wb["Records"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    meta = {k: v for k, v in wb["Metadata"].iter_rows(values_only=True)}
    return meta, [dict(zip(header, r, strict=True)) for r in body]


def pdf_text(content: bytes) -> str:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def pdf_ids(content: bytes) -> list[str]:
    """Ids from the records table in page order (the metadata block holds none)."""
    body = pdf_text(content).split("Document excerpts")[0]
    return UUID.findall(body)


def key(r: dict) -> str:
    return r["record_id"] or r["chunk_id"]
