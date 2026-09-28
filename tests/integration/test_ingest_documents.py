"""DOCX + text-PDF ingestion: wide tables, page tables, narrative chunks, cleaning."""

import io

import pytest
from docx import Document
from sqlalchemy import text

from app.db.session import owner_session
from tests.support.ingest import records, upload

pytestmark = pytest.mark.integration


def chunks(where="true", **params):
    with owner_session() as s:
        return [
            dict(r)
            for r in s.execute(
                text(f"SELECT * FROM document_chunks WHERE {where} ORDER BY locator"), params
            ).mappings()
        ]


def _by_locator_status(manifest):
    return {r["locator"]: r["status"] for r in manifest["rows"]}


def test_docx_wide_table(ingest_api, auth, manifests):
    body = upload(ingest_api, auth("a_eng_manager"), "tenant_a_week2.docx").json()
    assert body["status"] == "completed", body
    assert body["counts"]["records_created"] == 25 and body["counts"]["row_failures"] == 0
    recs = records("source_file = 'tenant_a_week2.docx'")
    assert {r["source_locator"]: r["status"] for r in recs} == _by_locator_status(
        manifests["files"]["tenant_a_week2.docx"]
    )
    assert {r["attendance_date"].isoformat() for r in recs} == {
        f"2026-09-{d:02d}" for d in range(7, 12)
    }
    assert all(r["extraction_method"] == "docx" for r in recs)


def test_docx_narrative_chunks(ingest_api, auth, manifests):
    body = upload(ingest_api, auth("a_eng_manager"), "tenant_a_week2.docx").json()
    assert body["counts"]["narrative_chunks"] == 5
    remarks = chunks("chunk_type = 'narrative' AND locator LIKE 'section=Manager remarks%'")
    expected = {
        n["locator"]: n["text"] for n in manifests["files"]["tenant_a_week2.docx"]["narrative"]
    }
    assert {c["locator"]: c["text"] for c in remarks} == {
        k: v for k, v in expected.items() if k.startswith("section=Manager remarks")
    }
    assert {c["classification"] for c in remarks} == {"confidential"}
    assert all(c["entity_id"] == "engineering" for c in remarks)
    assert all(c["embedding"] is not None for c in remarks)  # embedded since Step 9
    notes = chunks("locator = 'section=Notes;para=1'")
    assert notes[0]["classification"] == "internal"
    assert not chunks("text LIKE '%Acme Corp - Internal%'")  # page header excluded


def test_text_pdf_tables_and_cleaning(ingest_api, auth, manifests):
    body = upload(ingest_api, auth("b_manager"), "tenant_b_sep.pdf").json()
    assert body["status"] == "completed", body
    assert body["counts"]["records_created"] == 132
    recs = records("source_file = 'tenant_b_sep.pdf'")
    assert {r["source_locator"]: r["status"] for r in recs} == _by_locator_status(
        manifests["files"]["tenant_b_sep.pdf"]
    )
    assert {r["tenant_id"] for r in recs} == {"tenant_b"}
    assert any("repeated header/footer" in w for w in body["warnings"])
    narrative = chunks("source_document_id = :d", d=body["document_id"])
    assert narrative and not any(
        "Confidential" in c["text"] or c["text"].startswith("Page ") for c in narrative
    )


def test_scanned_pdf_routed_to_ocr(ingest_api, auth):
    body = upload(ingest_api, auth("a_hr_admin"), "scan_printed.pdf").json()
    assert body["status"] == "completed" and body["counts"]["records_created"] == 8
    assert any("printed sheet, engine=ocr_tesseract" in w for w in body["warnings"])
    assert {r["extraction_method"] for r in records("source_file = 'scan_printed.pdf'")} == {
        "ocr_tesseract"
    }


def test_conflict_detected_across_formats(ingest_api, auth, scope_for):
    upload(ingest_api, auth("a_hr_admin"), "tenant_a_sep_v2.csv")
    body = upload(ingest_api, auth("a_hr_admin"), "conflict_note.pdf").json()
    assert body["counts"]["records_created"] == 1
    from app.db.session import scoped_session

    with scoped_session(scope_for("a_hr_admin")) as s:
        row = s.execute(
            text(
                "SELECT status, reported_statuses, source_count FROM v_attendance "
                "WHERE employee_id = 'E002' AND attendance_date = '2026-09-15'"
            )
        ).one()
    assert row.status == "conflict" and list(row.reported_statuses) == ["absent", "present"]
    assert row.source_count == 2


def test_pii_in_narrative_masked(ingest_api, auth):
    doc = Document()
    doc.add_heading("HR follow-up", level=1)
    doc.add_paragraph(
        "Emma Wilson asked to be reached on +1-555-0107 or emma.w@example.com about leave."
    )
    buf = io.BytesIO()
    doc.save(buf)
    body = upload(ingest_api, auth("a_hr_admin"), "hr_note.docx", content=buf.getvalue()).json()
    assert body["status"] == "completed" and body["counts"]["pii_masked"] == 2
    (c,) = chunks("source_document_id = :d", d=body["document_id"])
    assert (
        "555-0107" not in c["text_masked"]
        and "[PHONE]" in c["text_masked"]
        and "[EMAIL]" in c["text_masked"]
    )
    assert "555-0107" in c["text"]  # raw kept for audit; rag_reader has no grant on it


def test_new_version_deactivates_old_chunks(ingest_api, auth):
    def note(txt):
        d = Document()
        d.add_paragraph(txt)
        b = io.BytesIO()
        d.save(b)
        return b.getvalue()

    v1 = upload(
        ingest_api, auth("a_hr_admin"), "policy_note.docx", content=note("Core hours start 09:30.")
    ).json()
    v2 = upload(
        ingest_api,
        auth("a_hr_admin"),
        "policy_note_v2.docx",
        content=note("Core hours start 09:00."),
    ).json()
    assert v2["version"] == 2
    assert [c["is_active"] for c in chunks("source_document_id = :d", d=v1["document_id"])] == [
        False
    ]
    assert [c["is_active"] for c in chunks("source_document_id = :d", d=v2["document_id"])] == [
        True
    ]
