"""Ingestion stage "indexed": row cards for every record, every active chunk embedded,
no PII in anything embedded, re-indexing on new versions, and outage recovery."""

import io

import pytest
from docx import Document
from sqlalchemy import text

from app.db.session import owner_session
from app.ingestion import service
from app.retrieval.embeddings import EmbeddingUnavailable, FakeEmbedder
from app.security.pii import contains_pii
from scripts.reindex import reindex
from tests.support.ingest import upload

pytestmark = pytest.mark.integration

TEXT_FILES = [
    ("a_hr_admin", "tenant_a_sep_v2.csv"),
    ("a_hr_admin", "tenant_a_sales_sep.xlsx"),
    ("a_eng_manager", "tenant_a_week2.docx"),
    ("a_hr_admin", "tenant_a_hr_contacts.xlsx"),
    ("b_manager", "tenant_b_sep.pdf"),
    ("a_hr_admin", "conflict_note.pdf"),
    ("a_eng_manager", "injection_memo.docx"),
]


def q(sql, **params):
    with owner_session() as s:
        return [dict(r) for r in s.execute(text(sql), params).mappings()]


def test_every_record_gets_one_row_card_and_every_chunk_is_embedded(ingest_api, auth):
    for persona, name in TEXT_FILES:
        body = upload(ingest_api, auth(persona), name).json()
        assert body["status"] == "completed", (name, body)
        assert body["counts"]["row_cards"] == body["counts"]["records_created"], name
        assert body["counts"]["embedding_pending"] == 0
        assert "indexed" in [h["stage"] for h in body["stage_history"]]
    cards = q(
        "SELECT r.record_id, c.chunk_id, c.tenant_id = r.tenant_id AS same_tenant, "
        "c.entity_id = r.entity_id AS same_entity, c.employee_id = r.employee_id AS same_emp, "
        "c.locator = r.source_locator AS same_loc "
        "FROM attendance_records r LEFT JOIN document_chunks c "
        "ON c.record_id = r.record_id AND c.chunk_type = 'row_card'"
    )
    assert cards and all(c["chunk_id"] for c in cards), "a record without a row card"
    assert len({c["record_id"] for c in cards}) == len(cards), "duplicate row cards"
    assert all(
        c["same_tenant"] and c["same_entity"] and c["same_emp"] and c["same_loc"] for c in cards
    )
    dims = q("SELECT vector_dims(embedding) AS d, count(*) AS n FROM document_chunks GROUP BY 1")
    assert dims == [{"d": 768, "n": dims[0]["n"]}], dims  # nothing NULL, all 768-dim


def test_row_card_text(ingest_api, auth):
    upload(ingest_api, auth("a_hr_admin"), "tenant_a_sep_v2.csv")
    card = q(
        "SELECT c.text_masked FROM document_chunks c JOIN attendance_records r USING (record_id) "
        "WHERE r.employee_id = 'E001' AND r.attendance_date = '2026-09-01'"
    )[0]["text_masked"]
    assert card.startswith("2026-09-01 | E001 Alice Johnson | Engineering | Present | ")
    assert card.endswith("| source tenant_a_sep_v2.csv row=2")


def test_review_records_get_cards_marked_review(ingest_api, auth, monkeypatch):
    from tests.support.vision import use_recorded_vision

    use_recorded_vision(monkeypatch, "handwritten_ambiguous")
    upload(ingest_api, auth("a_hr_admin"), "handwritten_ambiguous.png")
    rows = q(
        "SELECT r.review_required, c.text_masked FROM attendance_records r "
        "JOIN document_chunks c USING (record_id)"
    )
    assert rows and all(
        r["text_masked"].endswith("awaiting review") == r["review_required"] for r in rows
    )
    assert any(r["review_required"] for r in rows)


def test_no_pii_is_ever_embedded(ingest_api, auth, truth):
    doc = Document()
    doc.add_heading("Notes", level=1)
    doc.add_paragraph("Call Alice Johnson on +1-555-0100 or alice.johnson@tenant_a.example today.")
    buf = io.BytesIO()
    doc.save(buf)
    body = upload(
        ingest_api, auth("a_hr_admin"), "contact_note.docx", content=buf.getvalue()
    ).json()
    assert body["status"] == "completed", body
    for persona, name in TEXT_FILES:
        upload(ingest_api, auth(persona), name)
    embedded = q("SELECT text, text_masked FROM document_chunks WHERE embedding IS NOT NULL")
    assert embedded
    for c in embedded:
        assert not contains_pii(c["text_masked"]), c["text_masked"]
    note = next(c for c in embedded if "Call Alice" in c["text"])
    assert "[PHONE]" in note["text_masked"] and "[EMAIL]" in note["text_masked"]
    everything = " ".join(c["text_masked"] for c in embedded)
    for e in truth["employees"].values():
        assert e["phone"] not in everything and e["email"] not in everything
        assert e["national_id"] not in everything


def test_new_version_reindexes_and_deactivates_old_chunks(ingest_api, auth):
    v1 = upload(ingest_api, auth("a_hr_admin"), "tenant_a_sep.csv").json()
    v2 = upload(ingest_api, auth("a_hr_admin"), "tenant_a_sep_v2.csv").json()
    old = q(
        "SELECT is_active FROM document_chunks WHERE source_document_id = :d", d=v1["document_id"]
    )
    new = q(
        "SELECT is_active, embedding IS NOT NULL AS embedded FROM document_chunks "
        "WHERE source_document_id = :d",
        d=v2["document_id"],
    )
    assert old and not any(c["is_active"] for c in old)
    assert new and all(c["is_active"] and c["embedded"] for c in new)
    card = q(
        "SELECT c.text_masked FROM document_chunks c JOIN attendance_records r USING (record_id) "
        "WHERE c.is_active AND r.employee_id = 'E006' AND r.attendance_date = '2026-09-10'"
    )
    assert len(card) == 1 and "Present" in card[0]["text_masked"]  # the corrected row


def test_reprocessing_does_not_duplicate_chunks(ingest_api, auth, scope_for):
    body = upload(ingest_api, auth("a_hr_admin"), "tenant_a_week2.docx").json()
    n = q("SELECT count(*) AS n FROM document_chunks")[0]["n"]
    with owner_session() as s:
        s.execute(
            text("UPDATE ingestion_jobs SET status = 'queued' WHERE job_id = :j"),
            {"j": body["job_id"]},
        )
    assert (
        service.process(body["job_id"], service.scope_to_dict(scope_for("a_hr_admin")))
        == "completed"
    )
    assert q("SELECT count(*) AS n FROM document_chunks")[0]["n"] == n


class _Down:
    name, model, dim = "ollama", "nomic-embed-text", 768

    def embed_documents(self, texts):
        raise EmbeddingUnavailable("embedding backend unreachable: ConnectError")


def test_embedding_outage_defers_indexing_and_reindex_recovers(ingest_api, auth, monkeypatch):
    monkeypatch.setattr(service, "get_embedder", lambda: _Down())
    body = upload(ingest_api, auth("a_hr_admin"), "tenant_a_sep_v2.csv").json()
    assert body["status"] == "completed"  # data is stored and queryable
    assert body["counts"]["embedding_pending"] == 176
    assert "index_pending" in [h["stage"] for h in body["stage_history"]]
    assert any("embeddings pending for 176 chunks" in w for w in body["warnings"])
    assert q("SELECT count(*) AS n FROM document_chunks WHERE embedding IS NULL")[0]["n"] == 176

    result = reindex()
    assert result == {"row_cards": 176, "embedded": 176, "pending": 0, "model": FakeEmbedder.model}
    assert q("SELECT count(*) AS n FROM document_chunks WHERE embedding IS NULL")[0]["n"] == 0
    assert reindex()["embedded"] == 0  # idempotent


def test_reindex_reset_reembeds_everything(ingest_api, auth):
    upload(ingest_api, auth("a_eng_manager"), "tenant_a_week2.docx")
    n = q("SELECT count(*) AS n FROM document_chunks WHERE is_active")[0]["n"]
    assert reindex(reset=True)["embedded"] == n


def test_deep_health_reports_embeddings_and_vector_index(client):
    body = client.get("/v1/health/deep").json()["components"]
    assert body["embeddings"]["status"] == "ok"
    assert (body["embeddings"]["provider"], body["embeddings"]["dim"]) == ("fake", 768)
    assert body["vector"]["status"] == "ok" and "hnsw" in body["vector"]["engine"]


def test_deep_health_degraded_when_embedder_down(client, monkeypatch):
    from app.retrieval import embeddings

    class Unreachable(_Down):
        def ping(self):
            return {"reachable": False, "model_available": False}

    monkeypatch.setattr(embeddings, "get_embedder", lambda: Unreachable())
    body = client.get("/v1/health/deep").json()
    assert body["components"]["embeddings"]["status"] == "down"
    assert body["status"] == "degraded"  # core up, a model dependency down


def test_schema_vector_width_matches_the_embedder(corpus_db):
    from app.config import get_settings

    types = q(
        "SELECT attrelid::regclass::text AS t, format_type(atttypid, atttypmod) AS ty "
        "FROM pg_attribute WHERE attname IN ('embedding', 'question_embedding') "
        "AND attrelid IN ('document_chunks'::regclass, 'feedback_examples'::regclass)"
    )
    assert {r["ty"] for r in types} == {f"vector({get_settings().embed_dim})"} == {"vector(768)"}


def test_reindex_backfills_row_cards_for_older_data(ingest_api, auth):
    upload(ingest_api, auth("a_hr_admin"), "tenant_a_sep_v2.csv")
    with owner_session() as s:  # simulate data ingested before Step 9
        s.execute(text("DELETE FROM document_chunks WHERE chunk_type = 'row_card'"))
    result = reindex()
    assert result["embedded"] == 176 and result["pending"] == 0
    assert (
        q("SELECT count(*) AS n FROM document_chunks WHERE chunk_type = 'row_card'")[0]["n"] == 176
    )
