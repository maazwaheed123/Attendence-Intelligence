"""Every rendered file must actually contain what its manifest says, at the stated locator."""

import csv
import re

import pdfplumber
import pytest
from docx import Document
from openpyxl import load_workbook

pytestmark = pytest.mark.unit


def _loc(locator: str) -> dict:
    return dict(part.split("=", 1) for part in locator.split(";"))


def _truth_index(truth):
    return {(r["tenant_id"], r["employee_id"], r["attendance_date"]): r for r in truth["rows"]}


@pytest.mark.parametrize("name", ["tenant_a_sep.csv", "tenant_a_sep_v2.csv", "other_product.csv"])
def test_csv_rows_match_manifest(data_dir, manifests, name):
    with open(data_dir / "generated" / name, newline="", encoding="utf-8") as f:
        lines = list(csv.reader(f))
    header = lines[0]
    m = manifests["files"][name]
    assert len(lines) - 1 == len(m["rows"])
    for row in m["rows"]:
        values = dict(zip(header, lines[int(_loc(row["locator"])["row"]) - 1], strict=True))
        assert values == row["rendered"]
        assert values["employee_id"] == row["employee_id"]
        assert values["date"] == row["attendance_date"]


def test_v1_has_exactly_three_wrong_rows_and_v2_none(manifests):
    v1 = manifests["files"]["tenant_a_sep.csv"]["rows"]
    v2 = manifests["files"]["tenant_a_sep_v2.csv"]["rows"]
    assert sum(not r["matches_truth"] for r in v1) == 3
    assert all(r["matches_truth"] for r in v2)
    assert manifests["files"]["tenant_a_sep_v2.csv"]["logical_name"] == "tenant_a_sep"
    assert manifests["files"]["tenant_a_sep_v2.csv"]["version"] == 2


@pytest.mark.parametrize("name", ["tenant_a_sales_sep.xlsx", "tenant_a_hr_contacts.xlsx"])
def test_xlsx_cells_match_manifest(data_dir, manifests, name):
    wb = load_workbook(data_dir / "generated" / name)
    for row in manifests["files"][name]["rows"]:
        loc = _loc(row["locator"])
        ws = wb[loc["sheet"]]
        cells = [c.value for c in ws[int(loc["row"])]]
        assert cells[0] == row["employee_id"]
        status_value = row["rendered"].get("Attendance") or row["rendered"].get("Status")
        assert status_value in cells


def test_xlsx_lucas_gap(manifests):
    keys = {
        (r["employee_id"], r["attendance_date"])
        for r in manifests["files"]["tenant_a_sales_sep.xlsx"]["rows"]
    }
    assert ("E011", "2026-09-30") not in keys
    assert ("E010", "2026-09-30") in keys


def test_pii_file_contains_restricted_values(data_dir, truth):
    wb = load_workbook(data_dir / "generated" / "tenant_a_hr_contacts.xlsx")
    ws = wb["HR Export"]
    phones = {c.value for c in ws["D"][1:]}
    assert all(re.fullmatch(r"\+1-555-01\d\d", p) for p in phones)
    roster_phones = {e["phone"] for e in truth["employees"].values()}
    assert phones <= roster_phones


def test_docx_wide_table_and_remarks(data_dir, manifests, truth):
    doc = Document(data_dir / "generated" / "tenant_a_week2.docx")
    table = doc.tables[0]
    header = [c.text for c in table.rows[0].cells]
    idx = _truth_index(truth)
    for row in manifests["files"]["tenant_a_week2.docx"]["rows"]:
        loc = _loc(row["locator"])
        cells = table.rows[int(loc["row"]) - 1].cells
        assert cells[0].text == row["employee_id"]
        assert cells[header.index(loc["col"])].text == row["rendered"][loc["col"]]
        assert (
            idx[("tenant_a", row["employee_id"], row["attendance_date"])]["status"] == row["status"]
        )
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Manager remarks" in text
    assert "Bob Smith (E002) arrived late" in text
    assert "03/09/2026" in text
    assert doc.sections[0].header.paragraphs[0].text == "Acme Corp - Internal"


def test_docx_remarks_consistent_with_truth(truth):
    """The narrative claims must be true in the structured data."""
    idx = _truth_index(truth)
    bob = [idx[("tenant_a", "E002", f"2026-09-{d:02d}")] for d in range(7, 12)]
    assert [r["status"] for r in bob] == ["present", "present", "absent", "present", "absent"]
    assert all(r["check_in"] > "09:30" for r in bob if r["status"] == "present")
    assert idx[("tenant_a", "E001", "2026-09-03")]["status"] == "present"


def test_injection_memo_contains_payload(data_dir, manifests):
    text = "\n".join(
        p.text for p in Document(data_dir / "generated" / "injection_memo.docx").paragraphs
    )
    assert "Ignore all previous instructions" in text
    flagged = [n for n in manifests["files"]["injection_memo.docx"]["narrative"] if n["flags"]]
    assert len(flagged) == 1


def test_tenant_b_pdf_tables_match_manifest(data_dir, manifests):
    rows = manifests["files"]["tenant_b_sep.pdf"]["rows"]
    with pdfplumber.open(data_dir / "generated" / "tenant_b_sep.pdf") as pdf:
        assert len(pdf.pages) == 5
        tables = {i + 1: p.extract_tables()[0] for i, p in enumerate(pdf.pages)}
        for p in pdf.pages:
            text = p.extract_text()
            assert "Globex Ltd. - Confidential" in text and "Page " in text
    for row in rows:
        loc = _loc(row["locator"])
        table = tables[int(loc["page"])]
        cells = dict(zip(table[0], table[int(loc["row"]) - 1], strict=True))
        assert cells == row["rendered"]
    assert all(r["tenant_id"] == "tenant_b" for r in rows)


def test_conflict_note_contradicts_csv(data_dir, manifests):
    note = manifests["files"]["conflict_note.pdf"]["rows"][0]
    csv_row = next(
        r
        for r in manifests["files"]["tenant_a_sep_v2.csv"]["rows"]
        if (r["employee_id"], r["attendance_date"])
        == (note["employee_id"], note["attendance_date"])
    )
    assert note["status"] == "present" and csv_row["status"] == "absent"
    with pdfplumber.open(data_dir / "generated" / "conflict_note.pdf") as pdf:
        assert "was present at the client site on 15 September 2026" in pdf.pages[0].extract_text()


def test_every_truth_row_has_a_source(truth, manifests):
    covered = set()
    for m in manifests["files"].values():
        if m["kind"] == "attendance":
            covered |= {(r["tenant_id"], r["employee_id"], r["attendance_date"]) for r in m["rows"]}
    missing = [
        r
        for r in truth["rows"]
        if (r["tenant_id"], r["employee_id"], r["attendance_date"]) not in covered
    ]
    assert not missing
