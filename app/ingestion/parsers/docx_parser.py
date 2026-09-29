"""DOCX: tables (long or wide layout) -> records; paragraphs by heading -> narrative.

Page headers/footers are excluded (they are layout, not evidence).
"""

import io
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.ingestion.parsers.tabular import context_year, extract_rows, extract_wide
from app.ingestion.types import NarrativeBlock, ParseResult, PermanentError


def parse(content: bytes, filename: str) -> ParseResult:
    try:
        doc = Document(io.BytesIO(content))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise PermanentError(f"could not open document: {exc.__class__.__name__}") from exc

    result = ParseResult(method="docx")
    body_text = "\n".join(p.text for p in doc.paragraphs)
    year = context_year(body_text)

    for t_idx, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        long_ok = extract_rows(
            rows, result, lambda i, t=t_idx: f"table={t};row={i + 1}", source_label=f"table {t_idx}"
        )
        if not long_ok:
            wide_ok = extract_wide(
                rows,
                result,
                lambda i, col, t=t_idx: f"table={t};row={i + 1};col={col}",
                year=year,
                source_label=f"table {t_idx}",
            )
            if not wide_ok:
                result.warnings.append(f"table {t_idx}: not an attendance table, skipped")

    section, para_no = "Document", 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if (p.style.name or "").lower().startswith(("heading", "title")):
            section, para_no = text, 0
            continue
        para_no += 1
        result.narrative.append(
            NarrativeBlock(locator=f"section={section};para={para_no}", section=section, text=text)
        )

    if not result.records and not result.narrative:
        raise PermanentError("document contains no attendance table and no text")
    return result
