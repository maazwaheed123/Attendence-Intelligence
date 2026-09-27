import io
import zipfile

from openpyxl import load_workbook

from app.ingestion.parsers.tabular import extract_rows
from app.ingestion.types import ParseResult, PermanentError


def parse(content: bytes, filename: str) -> ParseResult:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise PermanentError(f"could not open workbook: {exc.__class__.__name__}") from exc

    result = ParseResult(method="xlsx")
    sheets_with_data = 0
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        title = ws.title
        # XLSX locator: sheet name + Excel row number (1-based), exactly what a user sees.
        if extract_rows(
            rows, result, lambda i, t=title: f"sheet={t};row={i + 1}", source_label=title
        ):
            sheets_with_data += 1
        elif any(any(c not in (None, "") for c in r) for r in rows):
            result.warnings.append(f"sheet '{title}': no attendance header found, skipped")
    if not sheets_with_data:
        raise PermanentError("no sheet contains an attendance table")
    return result
