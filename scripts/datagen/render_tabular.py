"""CSV and XLSX renderers."""

import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from scripts.datagen.manifest import STATUS_CODES, STATUS_WORDS, Manifest
from scripts.datagen.util import FIXED_DT, normalize_zip, to_date


def _flip(status: str) -> str:
    return "present" if status == "absent" else "absent"


def write_attendance_csv(path: Path, rows, *, logical_name, version, errors=frozenset()):
    """Long-format CSV. `errors` = keys rendered with a wrong status (v1 of the file)."""
    m = Manifest(
        path.name,
        "csv",
        rows[0]["tenant_id"],
        logical_name=logical_name,
        version=version,
        description="Engineering + HR, full month, long format",
    )
    header = [
        "date",
        "employee_id",
        "employee_name",
        "department",
        "status",
        "check_in",
        "check_out",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for i, r in enumerate(rows, start=2):
            wrong = (r["employee_id"], r["attendance_date"]) in errors
            st = _flip(r["status"]) if wrong else r["status"]
            cin = r["check_in"] if st == r["status"] else ("" if st == "absent" else "09:00")
            cout = r["check_out"] if st == r["status"] else ("" if st == "absent" else "17:30")
            vals = [
                r["attendance_date"],
                r["employee_id"],
                r["employee_name"],
                r["department"],
                STATUS_WORDS[st],
                cin or "",
                cout or "",
            ]
            w.writerow(vals)
            m.add_row(
                f"row={i}",
                r,
                status=st,
                rendered=dict(zip(header, vals, strict=True)),
                note="v1 error, corrected in v2" if wrong else None,
            )
    return m


def write_other_product_csv(path: Path, rows):
    m = Manifest(
        path.name,
        "csv",
        "tenant_a",
        product_id="hrms_ai",
        logical_name="hrms_training_attendance",
        description="Decoy product hrms_ai for product-isolation tests",
    )
    header = [
        "date",
        "employee_id",
        "employee_name",
        "department",
        "status",
        "check_in",
        "check_out",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for i, r in enumerate(rows, start=2):
            vals = [
                r["attendance_date"],
                r["employee_id"],
                r["employee_name"],
                r["department"],
                STATUS_WORDS[r["status"]],
                r["check_in"] or "",
                r["check_out"] or "",
            ]
            w.writerow(vals)
            m.add_row(f"row={i}", r, rendered=dict(zip(header, vals, strict=True)))
    return m


def _time_cell(s):
    return datetime.strptime(s, "%H:%M").time() if s else None


def _ampm(s):
    return datetime.strptime(s, "%H:%M").strftime("%I:%M %p").lstrip("0") if s else ""


def write_sales_xlsx(path: Path, rows, *, skip_keys=frozenset()):
    """Two sheets with DIFFERENT header names, title rows, merged cells, a totals row,
    real date/time cells on sheet 1 and text dates/12h times on sheet 2."""
    m = Manifest(
        path.name,
        "xlsx",
        "tenant_a",
        logical_name="tenant_a_sales_sep",
        description="Sales, full month, messy layout (one entry intentionally missing)",
    )
    wb = Workbook()
    wb.properties.creator = "HRIS export"
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT

    split = date(2026, 9, 15)
    part1 = [r for r in rows if to_date(r["attendance_date"]) <= split]
    part2 = [r for r in rows if to_date(r["attendance_date"]) > split]

    ws = wb.active
    ws.title = "Sep 1-15"
    ws["A1"] = "Acme Corp - Sales Department Attendance"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:G1")
    ws["A2"] = "Period: 01/09/2026 - 15/09/2026   Source: HRIS export"
    h1 = ["Emp Code", "Employee Name", "Dept", "Date", "Attendance", "In Time", "Out Time"]
    ws.append([])
    ws.append(h1)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    rownum = 4
    present_total = 0
    for r in part1:
        if (r["employee_id"], r["attendance_date"]) in skip_keys:
            continue
        rownum += 1
        vals = [
            r["employee_id"],
            r["employee_name"],
            r["department"],
            to_date(r["attendance_date"]),
            STATUS_CODES[r["status"]],
            _time_cell(r["check_in"]),
            _time_cell(r["check_out"]),
        ]
        ws.append(vals)
        ws.cell(rownum, 4).number_format = "DD/MM/YYYY"
        ws.cell(rownum, 6).number_format = "HH:MM"
        ws.cell(rownum, 7).number_format = "HH:MM"
        present_total += r["status"] == "present"
        m.add_row(
            f"sheet=Sep 1-15;row={rownum}",
            r,
            rendered={
                "Emp Code": vals[0],
                "Date": r["attendance_date"],
                "Attendance": vals[4],
                "In Time": r["check_in"] or "",
                "Out Time": r["check_out"] or "",
            },
        )
    ws.append([])
    ws.append(["Total present days", None, None, None, present_total])
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["D"].width = 12

    ws2 = wb.create_sheet("Sep 16-30")
    ws2["A1"] = "Sales Attendance (continued)"
    h2 = ["Staff ID", "Name", "Department", "Day", "Status", "Clock In", "Clock Out"]
    ws2.append(h2)
    rownum = 2
    for r in part2:
        if (r["employee_id"], r["attendance_date"]) in skip_keys:
            continue
        rownum += 1
        day = to_date(r["attendance_date"]).strftime("%d-%b-%Y")
        vals = [
            r["employee_id"],
            r["employee_name"],
            r["department"],
            day,
            STATUS_CODES[r["status"]],
            _ampm(r["check_in"]),
            _ampm(r["check_out"]),
        ]
        ws2.append(vals)
        m.add_row(
            f"sheet=Sep 16-30;row={rownum}",
            r,
            rendered=dict(zip(h2, vals, strict=True)),
        )
    wb.save(path)
    normalize_zip(path)
    return m


def write_pii_xlsx(path: Path, rows, employees):
    """HR attendance export that also carries restricted contact data (phone, national ID)."""
    m = Manifest(
        path.name,
        "xlsx",
        "tenant_a",
        logical_name="tenant_a_hr_contacts_attendance",
        description="HR, 1-4 Sep, includes restricted PII columns (fictional values)",
    )
    wb = Workbook()
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT
    ws = wb.active
    ws.title = "HR Export"
    header = [
        "Employee ID",
        "Full Name",
        "Department",
        "Phone",
        "National ID",
        "Email",
        "Date",
        "Status",
    ]
    ws.append(header)
    for i, r in enumerate(rows, start=2):
        e = employees[f"{r['tenant_id']}:{r['employee_id']}"]
        vals = [
            r["employee_id"],
            r["employee_name"],
            r["department"],
            e["phone"],
            e["national_id"],
            e["email"],
            to_date(r["attendance_date"]),
            STATUS_WORDS[r["status"]],
        ]
        ws.append(vals)
        ws.cell(i, 7).number_format = "YYYY-MM-DD"
        m.add_row(
            f"sheet=HR Export;row={i}",
            r,
            rendered={
                "Employee ID": vals[0],
                "Phone": vals[3],
                "National ID": vals[4],
                "Date": r["attendance_date"],
                "Status": vals[7],
            },
        )
    wb.save(path)
    normalize_zip(path)
    return m


def write_invalid_files(out: Path):
    corrupt = out / "corrupt.xlsx"
    corrupt.write_bytes(b"This is not a real spreadsheet.\x00\x01\x02 corrupted payload\n" * 4)
    empty = out / "empty.csv"
    empty.write_bytes(b"")
    return [
        Manifest(
            "corrupt.xlsx",
            "xlsx",
            "tenant_a",
            kind="invalid",
            expected_failure="not a valid XLSX/zip container",
            description="Corrupt file for validation and failure-reporting tests",
        ),
        Manifest(
            "empty.csv",
            "csv",
            "tenant_a",
            kind="invalid",
            expected_failure="empty file",
            description="Zero-byte file for validation tests",
        ),
    ]
